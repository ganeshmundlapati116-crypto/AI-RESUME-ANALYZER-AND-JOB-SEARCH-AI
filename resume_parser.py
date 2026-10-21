import pdfplumber
from docx import Document
from docx.shared import Inches, Pt
import re


# =========================================================
# EXTRACT RESUME TEXT
# =========================================================

def extract_text(file_path):

    text = ""

    if file_path.lower().endswith(".pdf"):

        with pdfplumber.open(file_path) as pdf:

            for page in pdf.pages:

                page_text = page.extract_text()

                if page_text:
                    text += page_text + "\n"

    elif file_path.lower().endswith(".docx"):

        doc = Document(file_path)

        for para in doc.paragraphs:
            text += para.text + "\n"

    return text


# =========================================================
# EXTRACT EMAIL
# =========================================================

def extract_email(text):

    email = re.findall(
        r'[\w\.-]+@[\w\.-]+\.\w+',
        text
    )

    if email:
        return email[0]

    return "Not Found"


# =========================================================
# EXTRACT PHONE
# =========================================================

def extract_phone(text):

    phone = re.findall(
        r'(?<!\d)(?:\+91[\s-]?)?[6-9]\d{9}(?!\d)',
        text
    )

    if phone:
        return phone[0]

    return "Not Found"


# =========================================================
# EXTRACT NAME
# =========================================================

def extract_name(text):

    lines = text.split("\n")

    for line in lines:

        line = line.strip()

        if line:
            return line

    return "Not Found"


# =========================================================
# NORMALIZE TEXT
# =========================================================

def normalize_text(text):

    text = text.lower()

    # Convert special separators into spaces
    text = re.sub(r'[/|,;:()\[\]{}]', ' ', text)

    # Remove extra spaces
    text = re.sub(r'\s+', ' ', text)

    return text.strip()


# =========================================================
# SKILL ALIASES
# =========================================================

SKILL_ALIASES = {

    "Python": [
        "python",
        "python programming",
        "python developer",
        "python language"
    ],

    "Java": [
        "java",
        "java programming",
        "java developer"
    ],

    "C": [
        "c programming",
        "c language"
    ],

    "C++": [
        "c++",
        "cpp",
        "c plus plus"
    ],

    "SQL": [
        "sql",
        "mysql",
        "postgresql",
        "database sql"
    ],

    "HTML": [
        "html",
        "html5"
    ],

    "CSS": [
        "css",
        "css3"
    ],

    "JavaScript": [
        "javascript",
        "java script",
        "js"
    ],

    "Flask": [
        "flask",
        "flask framework"
    ],

    "Django": [
        "django",
        "django framework"
    ],

    "React": [
        "react",
        "reactjs",
        "react.js"
    ],

    "Node.js": [
        "node.js",
        "nodejs",
        "node js"
    ],

    "Machine Learning": [
        "machine learning",
        "machine-learning",
        "ml"
    ],

    "Deep Learning": [
        "deep learning",
        "deep-learning",
        "dl"
    ],

    "TensorFlow": [
        "tensorflow",
        "tensor flow"
    ],

    "Keras": [
        "keras"
    ],

    "Pandas": [
        "pandas"
    ],

    "NumPy": [
        "numpy",
        "num py"
    ],

    "OpenCV": [
        "opencv",
        "open cv"
    ],

    "Git": [
        "git",
        "git version control"
    ],

    "GitHub": [
        "github",
        "git hub"
    ],

    "Excel": [
        "excel",
        "microsoft excel",
        "ms excel"
    ],

    "Power BI": [
        "power bi",
        "powerbi"
    ],

    "AWS": [
        "aws",
        "amazon web services"
    ],

    "Docker": [
        "docker",
        "docker container"
    ]
}


# =========================================================
# CHECK WHETHER SKILL EXISTS
# =========================================================

def skill_exists(text, skill):

    normalized = normalize_text(text)

    aliases = SKILL_ALIASES.get(skill, [skill.lower()])

    for alias in aliases:

        alias_normalized = normalize_text(alias)

        # Word-boundary based matching
        pattern = r'\b' + re.escape(alias_normalized) + r'\b'

        if re.search(pattern, normalized):
            return True

    return False


# =========================================================
# EXTRACT SKILLS
# =========================================================

def extract_skills(text):

    found_skills = []

    for skill in SKILL_ALIASES:

        if skill_exists(text, skill):

            found_skills.append(skill)

    return found_skills


# =========================================================
# EXTRACT EDUCATION
# =========================================================

def extract_education(text):

    education_keywords = [

        "B.Tech",
        "B. Tech",
        "BTech",
        "Bachelor of Technology",

        "B.E",
        "B.E.",
        "BE",
        "Bachelor of Engineering",

        "M.Tech",
        "M. Tech",
        "MTech",

        "MCA",
        "BCA",

        "Intermediate",

        "SSC",

        "Diploma"

    ]

    education = []

    lines = text.split("\n")

    for line in lines:

        clean_line = line.strip()

        if not clean_line:
            continue

        for keyword in education_keywords:

            if keyword.lower() in clean_line.lower():

                if clean_line not in education:
                    education.append(clean_line)

                break

    return education



# =========================================================
# EXTRACT PROJECTS - CLEAN VERSION
# =========================================================

def extract_projects(text):

    lines = text.splitlines()

    projects = []

    project_headings = [
        "project",
        "projects",
        "project experience",
        "project experience:",
        "academic projects",
        "major project",
        "minor project"
    ]

    stop_headings = [
        "education",
        "educational qualification",
        "skills",
        "technical skills",
        "experience",
        "work experience",
        "internship",
        "certifications",
        "certificate",
        "achievements",
        "professional summary",
        "career objective",
        "objective",
        "personal details",
        "declaration",
        "languages",
        "interests",
        "hobbies",
        "references"
    ]

    inside_project_section = False

    for line in lines:

        clean_line = line.strip()

        if not clean_line:
            continue

        lower_line = clean_line.lower()

        # START PROJECT SECTION
        if lower_line in project_headings:

            inside_project_section = True
            continue

        # STOP PROJECT SECTION
        if inside_project_section:

            if lower_line in stop_headings:
                break

            if len(clean_line) < 3:
                continue

            # Ignore common unwanted lines
            unwanted = [
                "project experience",
                "projects",
                "project"
            ]

            if lower_line in unwanted:
                continue

            # -------------------------------------------------
            # Split multiple projects when PDF extraction
            # puts them on the same line.
            # -------------------------------------------------

            parts = re.split(
                r'(?=(?:AI-Based Weather Prediction System|'
                r'File Encryption System|'
                r'UPI Fraud Detection))',
                clean_line,
                flags=re.IGNORECASE
            )

            for part in parts:

                part = part.strip()

                if not part:
                    continue

                if part not in projects:

                    projects.append(part)

    return projects


# =========================================================
# JOB ROLE SKILLS
# =========================================================

ROLE_SKILLS = {

    "Software Developer": [

        "Python",
        "Java",
        "SQL",
        "HTML",
        "CSS",
        "JavaScript",
        "Git"

    ],

    "Python Developer": [

        "Python",
        "Flask",
        "Django",
        "SQL",
        "Git"

    ],

    "AI Engineer": [

        "Python",
        "Machine Learning",
        "Deep Learning",
        "TensorFlow",
        "Pandas",
        "NumPy"

    ],

    "Data Analyst": [

        "Python",
        "SQL",
        "Pandas",
        "NumPy",
        "Excel"

    ],

    "Full Stack Developer": [

        "HTML",
        "CSS",
        "JavaScript",
        "React",
        "Node.js",
        "SQL",
        "Git"

    ]

}


# =========================================================
# FIND MISSING SKILLS
# =========================================================

def find_missing_skills(role, skills):

    if not role:
        role = "Software Developer"

    required_skills = ROLE_SKILLS.get(
        role,
        []
    )

    # Normalize skills coming from extract_skills()
    user_skills = set(
        skill.lower().strip()
        for skill in skills
    )

    missing = []

    for required_skill in required_skills:

        if required_skill.lower() not in user_skills:

            missing.append(required_skill)

    return missing


# ==========================
# PROFESSIONAL ATS SCORE
# ==========================

def calculate_ats_score(
    name,
    email,
    phone,
    skills,
    education,
    projects,
    resume_text="",
    role="Software Developer"
):

    resume_lower = resume_text.lower()

    # ==========================
    # 1. Contact Information - 10
    # ==========================

    contact_score = 0

    if name != "Not Found":
        contact_score += 3

    if email != "Not Found":
        contact_score += 3

    if phone != "Not Found":
        contact_score += 4


    # ==========================
    # 2. General Skills - 20
    # ==========================

    skill_score = min(
        len(skills) * 2,
        20
    )


    # ==========================
    # 3. Job Role Keyword Match - 25
    # ==========================

    role_skills = {

        "Software Developer": [
            "Python",
            "Java",
            "SQL",
            "HTML",
            "CSS",
            "JavaScript",
            "Git"
        ],

        "Python Developer": [
            "Python",
            "Flask",
            "Django",
            "SQL",
            "Git"
        ],

        "AI Engineer": [
            "Python",
            "Machine Learning",
            "Deep Learning",
            "TensorFlow",
            "Pandas",
            "NumPy"
        ],

        "Data Analyst": [
            "Python",
            "SQL",
            "Pandas",
            "NumPy",
            "Excel"
        ],

        "Full Stack Developer": [
            "HTML",
            "CSS",
            "JavaScript",
            "React",
            "Node.js",
            "SQL",
            "Git"
        ]
    }

    required_skills = role_skills.get(
        role,
        []
    )

    matched_skills = 0

    for required_skill in required_skills:

        if required_skill.lower() in resume_lower:
            matched_skills += 1

    if len(required_skills) > 0:

        keyword_score = round(
            (matched_skills / len(required_skills)) * 25
        )

    else:

        keyword_score = 0


    # ==========================
    # 4. Education - 15
    # ==========================

    education_score = 0

    if len(education) > 0:
        education_score = 15


    # ==========================
    # 5. Projects - 15
    # ==========================

    project_score = 0

    if len(projects) >= 2:

        project_score = 15

    elif len(projects) == 1:

        project_score = 10


    # ==========================
    # 6. Resume Content - 15
    # ==========================

    content_keywords = [
        "experience",
        "skills",
        "education",
        "project",
        "certification",
        "achievement",
        "summary",
        "objective"
    ]

    content_matches = 0

    for keyword in content_keywords:

        if keyword in resume_lower:
            content_matches += 1

    content_score = min(
        content_matches * 2,
        15
    )


    # ==========================
    # FINAL SCORE
    # ==========================

    total_score = (
        contact_score
        + skill_score
        + keyword_score
        + education_score
        + project_score
        + content_score
    )

    total_score = min(
        total_score,
        100
    )


    # Return all scores
    return {
        "total": total_score,
        "contact": contact_score,
        "skills": skill_score,
        "keyword": keyword_score,
        "education": education_score,
        "projects": project_score,
        "content": content_score
    }


# =========================================================
# KEYWORD MATCH PERCENTAGE
# =========================================================

def calculate_keyword_match(
    resume_text,
    role
):

    if not role:
        role = "Software Developer"

    required_skills = ROLE_SKILLS.get(
        role,
        []
    )

    if not required_skills:
        return 0

    matched = 0

    for skill in required_skills:

        if skill_exists(
            resume_text,
            skill
        ):

            matched += 1


    percentage = (
        matched
        /
        len(required_skills)
    ) * 100

    return round(
        percentage
    )


# =========================================================
# RESUME SUGGESTIONS
# =========================================================

def generate_suggestions(
    email,
    phone,
    skills,
    education,
    projects
):

    suggestions = []


    if email == "Not Found":

        suggestions.append(
            "Add a professional email address."
        )


    if phone == "Not Found":

        suggestions.append(
            "Add your phone number."
        )


    if len(skills) < 5:

        suggestions.append(
            "Add more relevant technical skills."
        )


    if len(projects) == 0:

        suggestions.append(
            "Include at least one project."
        )


    if len(education) == 0:

        suggestions.append(
            "Add your education details."
        )


    if len(suggestions) == 0:

        suggestions.append(
            "Your resume has good basic ATS elements. Consider tailoring keywords to the selected job role."
        )


    return suggestions

# =========================================================
# GROUP PROJECT LINES
# =========================================================

def group_project_entries(projects):

    grouped_projects = []

    current_project = ""

    for project in projects:

        project = project.strip()

        if not project:
            continue

        # A line containing "|" is treated as a new project title
        if "|" in project:

            if current_project:

                grouped_projects.append(
                    current_project.strip()
                )

            current_project = project

        else:

            if current_project:

                current_project += " " + project

            else:

                current_project = project

    # Add final project
    if current_project:

        grouped_projects.append(
            current_project.strip()
        )

    return grouped_projects

# =========================================================
# CLEAN PROJECT DESCRIPTION
# =========================================================

def enhance_project_description(project):

    project = project.strip()

    if not project:
        return ""

    # Remove unwanted generic text
    project = re.sub(
        r'\bApplying relevant technical skills to develop a practical software solution\.?',
        '',
        project,
        flags=re.IGNORECASE
    )

    project = re.sub(
        r'\busing relevant technical skills\b',
        '',
        project,
        flags=re.IGNORECASE
    )

    # Improve common wording
    replacements = {
        "built a": "Developed a",
        "build a": "Developed a",
        "made a": "Developed a",
        "created a": "Developed a",
        "worked on": "Developed",
        "working on": "Developed",
        "used": "Utilized",
        "implemented": "Implemented"
    }

    for old, new in replacements.items():

        project = re.sub(
            r'\b' + re.escape(old) + r'\b',
            new,
            project,
            flags=re.IGNORECASE
        )

    # Remove duplicate spaces
    project = re.sub(
        r'\s+',
        ' ',
        project
    ).strip()

    return project

def generate_enhanced_resume(
    output_path,
    name,
    email,
    phone,
    skills,
    education,
    projects,
    role="Software Developer"
):

    doc = Document()

    # =====================================================
    # RESUME PAGE FORMATTING
    # =====================================================

    section = doc.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles

    normal_style = styles["Normal"]

    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9.5)

    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.line_spacing = 1.0

    heading_style = styles["Heading 1"]

    heading_style.font.name = "Arial"
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True

    heading_style.paragraph_format.space_before = Pt(7)
    heading_style.paragraph_format.space_after = Pt(3)


    # =====================================================
    # ROLE BASED KEYWORDS
    # =====================================================

    role_keywords = {

        "software developer": [
            "python",
            "java",
            "sql",
            "javascript",
            "html",
            "css",
            "react",
            "node.js",
            "git"
        ],

        "python developer": [
            "python",
            "flask",
            "django",
            "sql",
            "machine learning",
            "git"
        ],

        "web developer": [
            "html",
            "css",
            "javascript",
            "react",
            "node.js",
            "flask",
            "sql"
        ],

        "frontend developer": [
            "html",
            "css",
            "javascript",
            "react"
        ],

        "backend developer": [
            "python",
            "flask",
            "node.js",
            "sql",
            "git"
        ],

        "machine learning engineer": [
            "python",
            "machine learning",
            "sql",
            "git"
        ],

        "data scientist": [
            "python",
            "machine learning",
            "sql"
        ],

        "full stack developer": [
            "html",
            "css",
            "javascript",
            "react",
            "node.js",
            "python",
            "flask",
            "sql"
        ]
    }


    selected_role = role.strip().lower()

    target_keywords = role_keywords.get(
        selected_role,
        role_keywords["software developer"]
    )


    # =====================================================
    # CLEAN SKILLS
    # =====================================================

    unique_skills = []

    for skill in skills:

        skill = skill.strip()

        if not skill:
            continue

        if skill.lower() not in [
            s.lower() for s in unique_skills
        ]:

            unique_skills.append(skill)


    # =====================================================
    # ROLE-RELEVANT SKILLS FIRST
    # =====================================================

    relevant_skills = []
    other_skills = []

    for skill in unique_skills:

        skill_lower = skill.lower()

        matched = False

        for keyword in target_keywords:

            if keyword in skill_lower:

                matched = True
                break

        if matched:

            relevant_skills.append(skill)

        else:

            other_skills.append(skill)


    ordered_skills = (
        relevant_skills +
        other_skills
    )


    # =====================================================
    # CLEAN PROJECT DATA
    # =====================================================

    clean_projects = []

    unwanted_project_text = [
        "projects",
        "project experience",
        "professional summary",
        "career objective",
        "motivated software developer candidate",
        "to obtain an entry-level",
        "technical skills"
    ]

    for project in projects:

        project = project.strip()

        if not project:
            continue

        project_lower = project.lower()

        if any(
            text in project_lower
            for text in unwanted_project_text
        ):
            continue

        if project not in clean_projects:

            clean_projects.append(project)


    # =====================================================
    # ROLE-RELEVANT PROJECTS FIRST
    # =====================================================

    relevant_projects = []
    other_projects = []

    for project in clean_projects:

        project_lower = project.lower()

        matched = False

        for keyword in target_keywords:

            if keyword in project_lower:

                matched = True
                break

        if matched:

            relevant_projects.append(project)

        else:

            other_projects.append(project)


    ordered_projects = (
        relevant_projects +
        other_projects
    )


    # =====================================================
    # NAME
    # =====================================================

    doc.add_heading(
        name.upper(),
        level=0
    )


    # =====================================================
    # CONTACT
    # =====================================================

    contact_parts = []

    if email:
        contact_parts.append(email)

    if phone:
        contact_parts.append(phone)

    doc.add_paragraph(
        " | ".join(contact_parts)
    )


    # =====================================================
    # PROFESSIONAL SUMMARY
    # =====================================================

    doc.add_heading(
        "Professional Summary",
        level=1
    )

    if relevant_skills:

        summary_skills = relevant_skills[:6]

    else:

        summary_skills = unique_skills[:6]


    if summary_skills:

        skill_text = ", ".join(
            summary_skills
        )

        summary = (
            f"Motivated {role} candidate with a strong foundation "
            f"in software development and hands-on experience with "
            f"{skill_text}. Demonstrates practical knowledge through "
            f"academic and software projects and is interested in "
            f"applying technical skills to real-world development "
            f"solutions."
        )

    else:

        summary = (
            f"Motivated {role} candidate with a strong foundation "
            f"in software development and a passion for building "
            f"practical technology solutions."
        )


    doc.add_paragraph(
        summary
    )


    # =====================================================
    # TECHNICAL SKILLS
    # =====================================================

    doc.add_heading(
        "Technical Skills",
        level=1
    )

    if ordered_skills:

        doc.add_paragraph(
            ", ".join(ordered_skills)
        )

    else:

        doc.add_paragraph(
            "No technical skills identified."
        )


    # =====================================================
    # EDUCATION
    # =====================================================

    doc.add_heading(
        "Education",
        level=1
    )

    if education:

        for edu in education:

            edu = edu.strip()

            if edu:

                doc.add_paragraph(
                    edu,
                    style="List Bullet"
                )

    else:

        doc.add_paragraph(
            "Education details not available."
        )


    # =====================================================
    # PROJECTS
    # =====================================================

    doc.add_heading(
        "Projects",
        level=1
    )

    if ordered_projects:

        for project in ordered_projects:

            enhanced_project = enhance_project_description(
                project
            )

            if not enhanced_project:
                continue


            # =================================================
            # PROJECT TITLE + DESCRIPTION
            # =================================================

            if "|" in enhanced_project:

                parts = enhanced_project.split(
                    "|",
                    1
                )

                project_title = parts[0].strip()

                project_description = parts[1].strip()


                # Project title

                p_title = doc.add_paragraph()

                run = p_title.add_run(
                    project_title
                )

                run.bold = True


                # Project description

                if project_description:

                    doc.add_paragraph(
                        project_description,
                        style="List Bullet"
                    )

            else:

                doc.add_paragraph(
                    enhanced_project,
                    style="List Bullet"
                )

    else:

        doc.add_paragraph(
            "No project details were identified from the uploaded resume."
        )


    # =====================================================
    # CAREER OBJECTIVE
    # =====================================================

    doc.add_heading(
        "Career Objective",
        level=1
    )

    objective = (
        f"To obtain an entry-level {role} position where I can "
        f"apply my technical skills, develop practical expertise, "
        f"and contribute to professional software projects."
    )

    doc.add_paragraph(
        objective
    )


    # =====================================================
    # SAVE ENHANCED RESUME
    # =====================================================

    doc.save(
        output_path
    )

    return output_path